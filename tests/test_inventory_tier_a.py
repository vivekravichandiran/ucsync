"""Task 4: inventory completeness — Tier-A AI assets are discovered & reported
(report-only, in_scope_for_migration=false); Tier-B/C get no sheet; streaming tables
& materialized views are report-only (not migrated) unless opted in."""

from __future__ import annotations

import json
from pathlib import Path

from uc_sync.config import from_sources
from uc_sync.inventory import InventoryService
from uc_sync.models import ObjectType
from uc_sync.package_import import PackageImportEngine
from uc_sync.report import build_report


class PathAwareSource:
    """Fake UC REST source that dispatches by path — including the Tier-A
    registered-models list endpoint."""

    def paginate(self, path, items_key, **q):
        if path.endswith("/catalogs"):
            return iter([{"name": "c", "catalog_type": "MANAGED_CATALOG"}])
        if path.endswith("/schemas"):
            return iter([{"name": "s", "full_name": "c.s"}])
        if path.endswith("/tables"):
            return iter([])
        if path.endswith("/volumes"):
            return iter([])
        if path.endswith("/functions"):
            return iter([])
        if path.endswith("/external-locations"):
            return iter([])
        if path.endswith("/storage-credentials"):
            return iter([])
        if path.endswith("/models"):
            return iter([{"name": "fraud_model", "full_name": "c.s.fraud_model",
                          "owner": "ml@x.com", "comment": "scoring model"}])
        if path.endswith("/versions"):
            return iter([{"version": 1}, {"version": 2}])
        return iter([])

    def get(self, path, **q):
        return {}


def _cfg():
    return from_sources({
        "execution_mode": "LOCAL", "stage": "INVENTORY",
        "catalogs": "c", "output_volume_path": "/Volumes/o/u/v",
        "ops_catalog": "o", "ops_schema": "u",
    })


def test_registered_model_discovered_as_report_only():
    objects = InventoryService(PathAwareSource(), _cfg()).run()
    models = [o for o in objects if o.object_type == ObjectType.MODEL]
    assert len(models) == 1
    model = models[0]
    assert model.full_name == "c.s.fraud_model"
    assert model.definition.get("in_scope_for_migration") is False
    assert model.definition.get("version_count") == 2


def test_report_shows_tier_a_tab_and_no_tier_b_c_tabs(tmp_path: Path):
    objects = [
        {"object_type": "MODEL", "full_name": "c.s.m", "owner": "ml@x.com",
         "tags": {}, "grants": [],
         "definition": {"in_scope_for_migration": False, "comment": "m"}},
        {"object_type": "TABLE", "full_name": "c.s.t", "owner": "me",
         "tags": {}, "grants": [], "definition": {}},
    ]
    out = tmp_path / "r.xlsx"
    build_report(objects, str(out), stage="INVENTORY")
    from openpyxl import load_workbook
    wb = load_workbook(out)
    # Tier-A Models tab present with an in_scope_for_migration column reading false.
    assert "Models" in wb.sheetnames
    hdr = list(wb["Models"].iter_rows(values_only=True))[0]
    assert "in_scope_for_migration" in hdr
    row = list(wb["Models"].iter_rows(values_only=True))[1]
    assert row[hdr.index("in_scope_for_migration")] == "false"
    # No Tier-B / Tier-C sheets exist any more.
    for gone in ("Connections", "Service Credentials", "Foreign Catalogs",
                 "Shares", "Recipients", "Providers"):
        assert gone not in wb.sheetnames


def _write(root: Path, rel: str, body: str) -> None:
    p = root / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(body, encoding="utf-8")


class RecordingSql:
    def __init__(self):
        self.statements = []

    def execute(self, sql):
        self.statements.append(sql)
        return None


def _pipeline_bundle(tmp_path: Path) -> Path:
    root = tmp_path / "migrated"
    _write(root, "ddl/STREAMING_TABLE_c__s__st.sql",
           "CREATE STREAMING TABLE `c`.`s`.`st` AS SELECT * FROM stream;\n")
    _write(root, "ddl/MATERIALIZED_VIEW_c__s__mv.sql",
           "CREATE MATERIALIZED VIEW `c`.`s`.`mv` AS SELECT 1;\n")
    _write(root, "inventory/objects.json", "[]")
    return root


def test_streaming_table_and_mv_report_only_by_default(tmp_path: Path):
    root = _pipeline_bundle(tmp_path)
    sql = RecordingSql()
    results = PackageImportEngine(str(root), sql, dry_run=False).run()
    by = {r.object_type: r for r in results}
    assert by["STREAMING_TABLE"].action == "REPORT_ONLY"
    assert by["MATERIALIZED_VIEW"].action == "REPORT_ONLY"
    # Neither was created (no pipeline spun).
    assert not any("STREAMING TABLE" in s.upper() for s in sql.statements)
    assert not any("MATERIALIZED VIEW" in s.upper() for s in sql.statements)


def test_mv_migrated_when_opted_in(tmp_path: Path):
    root = _pipeline_bundle(tmp_path)
    sql = RecordingSql()
    results = PackageImportEngine(
        str(root), sql, dry_run=False, migrate_materialized_views=True
    ).run()
    by = {r.object_type: r for r in results}
    # Streaming table is ALWAYS report-only; the MV is now created.
    assert by["STREAMING_TABLE"].action == "REPORT_ONLY"
    assert by["MATERIALIZED_VIEW"].action != "REPORT_ONLY"
    assert any("MATERIALIZED VIEW" in s.upper() for s in sql.statements)
