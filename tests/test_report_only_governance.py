"""Governance on report-only objects (MVs / streaming tables) is skipped, not failed.

A materialized view (without ``migrate_materialized_views``) or a streaming table is
never created on the target. Its column tags, masks and a table-level ABAC policy on
it used to run anyway, fail against the missing object, and hard-fail the import as
a governance failure. They are now skipped; catalog/schema ABAC policies and the
governance of migrated tables still apply.
"""

from __future__ import annotations

import json
from pathlib import Path

from uc_sync.export import _safe_filename
from uc_sync.package_import import PackageImportEngine, governance_failures

_MISSING_ON_TARGET = ("`mv`", "`st`")


class _TargetSql:
    """Fails any statement on an object that does not exist on the target."""

    def __init__(self, missing=_MISSING_ON_TARGET):
        self.missing = missing
        self.statements: list[str] = []

    def execute(self, statement):
        self.statements.append(statement)
        if any(name in statement for name in self.missing):
            raise RuntimeError("[TABLE_OR_VIEW_NOT_FOUND] not found on target")
        return []


def _write(root: Path, folder: str, object_type: str, full_name: str, sql: str):
    path = root / folder / f"{_safe_filename(object_type, full_name)}.sql"
    path.write_text(sql, encoding="utf-8")


def _bundle(tmp_path: Path) -> Path:
    root = tmp_path / "migrated"
    for d in ("ddl", "tags", "policies", "abac", "inventory"):
        (root / d).mkdir(parents=True)
    _write(root, "ddl", "CATALOG", "c", "CREATE CATALOG IF NOT EXISTS `c`;\n")
    _write(root, "ddl", "SCHEMA", "c.s", "CREATE SCHEMA IF NOT EXISTS `c`.`s`;\n")
    _write(root, "ddl", "TABLE", "c.s.t",
           "CREATE TABLE IF NOT EXISTS `c`.`s`.`t` (ssn STRING);\n")
    _write(root, "ddl", "MATERIALIZED_VIEW", "c.s.mv",
           "CREATE MATERIALIZED VIEW `c`.`s`.`mv` AS SELECT ssn FROM `c`.`s`.`t`;\n")
    _write(root, "ddl", "STREAMING_TABLE", "c.s.st",
           "CREATE STREAMING TABLE `c`.`s`.`st` (ssn STRING);\n")
    # Tags on the migrated table and on both report-only objects.
    _write(root, "tags", "TABLE", "c.s.t",
           "ALTER TABLE `c`.`s`.`t` ALTER COLUMN ssn SET TAGS ('pii' = 'ssn');\n")
    _write(root, "tags", "MATERIALIZED_VIEW", "c.s.mv",
           "ALTER MATERIALIZED VIEW `c`.`s`.`mv` ALTER COLUMN ssn "
           "SET TAGS ('pii' = 'ssn');\n")
    _write(root, "tags", "STREAMING_TABLE", "c.s.st",
           "ALTER STREAMING TABLE `c`.`s`.`st` ALTER COLUMN ssn "
           "SET TAGS ('pii' = 'ssn');\n")
    # A classic mask on the MV.
    _write(root, "policies", "MATERIALIZED_VIEW", "c.s.mv",
           "ALTER MATERIALIZED VIEW `c`.`s`.`mv` ALTER COLUMN ssn "
           "SET MASK `c`.`s`.`mask_ssn`;\n")
    # ABAC: one policy ON the MV itself, one at catalog level (masks tagged columns).
    policies = [
        ("c.s.mv#policy:mv_mask", "TABLE", "c.s.mv",
         "CREATE POLICY mv_mask ON TABLE `c`.`s`.`mv` COLUMN MASK `c`.`s`.`mask_ssn` "
         "TO `account users` FOR TABLES MATCH COLUMNS has_tag('pii') AS x ON COLUMN x;\n"),
        ("c#policy:pii_mask", "CATALOG", "c",
         "CREATE POLICY pii_mask ON CATALOG `c` COLUMN MASK `c`.`s`.`mask_ssn` "
         "TO `account users` FOR TABLES MATCH COLUMNS has_tag('pii') AS x ON COLUMN x;\n"),
    ]
    rows = [
        {"object_type": "CATALOG", "full_name": "c"},
        {"object_type": "SCHEMA", "full_name": "c.s"},
        {"object_type": "TABLE", "full_name": "c.s.t"},
        {"object_type": "MATERIALIZED_VIEW", "full_name": "c.s.mv"},
        {"object_type": "STREAMING_TABLE", "full_name": "c.s.st"},
    ]
    for full_name, on_type, on_securable, sql in policies:
        _write(root, "abac", "ABAC_POLICY", full_name, sql)
        rows.append({
            "object_type": "ABAC_POLICY", "full_name": full_name,
            "definition": {"on_securable_type": on_type, "on_securable": on_securable},
        })
    (root / "inventory" / "objects.json").write_text(json.dumps(rows), encoding="utf-8")
    return root


def test_report_only_governance_is_skipped_not_a_governance_failure(tmp_path: Path):
    sql = _TargetSql()
    results = PackageImportEngine(
        str(_bundle(tmp_path)), sql, dry_run=False, abac_sql_executor=sql,
    ).run()

    assert governance_failures(results) == []
    assert not [r for r in results if r.status == "FAILURE"], [
        (r.object_type, r.target_full_name, r.message) for r in results
        if r.status == "FAILURE"
    ]
    # Nothing was ever sent against the objects that are not on the target.
    assert not [s for s in sql.statements if any(n in s for n in _MISSING_ON_TARGET)]
    by_name = {
        (r.object_type, r.target_full_name): r for r in results
        if r.action == "REPORT_ONLY"
    }
    assert ("MATERIALIZED_VIEW", "c.s.mv") in by_name
    assert ("STREAMING_TABLE", "c.s.st") in by_name
    assert "not applied" in by_name[("MATERIALIZED_VIEW", "c.s.mv")].message
    # The migrated table's governance and the catalog-level policy still apply.
    executed = "\n".join(sql.statements)
    assert "ALTER TABLE `c`.`s`.`t` ALTER COLUMN ssn SET TAGS" in executed
    assert "CREATE POLICY pii_mask ON CATALOG `c`" in executed


def test_migrated_mv_keeps_its_governance(tmp_path: Path):
    """With migrate_materialized_views=true the MV is created, so its governance runs."""
    sql = _TargetSql(missing=("`st`",))
    PackageImportEngine(
        str(_bundle(tmp_path)), sql, dry_run=False, abac_sql_executor=sql,
        migrate_materialized_views=True,
    ).run()
    assert any("MATERIALIZED VIEW `c`.`s`.`mv`" in s and "SET TAGS" in s
               for s in sql.statements)


# ---------------------------------------------------------------------------
# migrate_materialized_views=true: a fresh MV is dropped fail-closed like a table
# ---------------------------------------------------------------------------

import re  # noqa: E402


class _StatefulTarget:
    """Objects exist only once created (or if pre-seeded). Optionally fails the MV's
    tag assignment or the catalog ABAC policy."""

    _CREATE = re.compile(
        r"CREATE (?:OR REPLACE )?(?:TABLE|MATERIALIZED VIEW|STREAMING TABLE)"
        r"(?: IF NOT EXISTS)? `c`\.`s`\.`(\w+)`"
    )

    def __init__(self, existing=(), fail_mv_tag=False, fail_catalog_policy=False):
        self.exists = set(existing)
        self.fail_mv_tag = fail_mv_tag
        self.fail_catalog_policy = fail_catalog_policy
        self.statements: list[str] = []

    def execute(self, statement):
        self.statements.append(statement)
        created = self._CREATE.search(statement)
        if created:
            self.exists.add(created.group(1))
            return []
        if statement.startswith("DROP "):
            self.exists.discard(statement.rsplit("`", 2)[-2])
            return []
        if self.fail_catalog_policy and "ON CATALOG" in statement:
            raise RuntimeError("PERMISSION_DENIED: cannot create policy")
        if self.fail_mv_tag and "`mv`" in statement and "SET TAGS" in statement:
            raise RuntimeError("PERMISSION_DENIED: ASSIGN on governed tag pii")
        for name in ("t", "mv", "st"):
            if f"`c`.`s`.`{name}`" in statement and name not in self.exists:
                raise RuntimeError(f"[TABLE_OR_VIEW_NOT_FOUND] c.s.{name}")
        return []


def _run_migrating_mvs(tmp_path: Path, sql: _StatefulTarget, mv_level_policy=False):
    root = _bundle(tmp_path)
    if not mv_level_policy:
        # Catalog-level ABAC only. A policy attached ON the MV itself is created
        # before the MV exists (known ordering, out of scope) — see the last test.
        (root / "abac" / f"{_safe_filename('ABAC_POLICY', 'c.s.mv#policy:mv_mask')}.sql"
         ).unlink()
    return PackageImportEngine(
        str(root), sql, dry_run=False, abac_sql_executor=sql,
        migrate_materialized_views=True,
    ).run()


def _create_row(results, name):
    return next(r for r in results if r.target_full_name == name
                and r.object_type in ("TABLE", "MATERIALIZED_VIEW")
                and r.action not in ("APPLY_TAGS", "CREATE_POLICY"))


def test_fresh_mv_dropped_when_its_tag_fails(tmp_path: Path):
    sql = _StatefulTarget(fail_mv_tag=True)
    results = _run_migrating_mvs(tmp_path, sql)
    mv = _create_row(results, "c.s.mv")
    assert mv.action == "DROP_PROTECTION_FAILED"
    assert mv.error_code == "PROTECTION_FAILED"
    assert "DROP MATERIALIZED VIEW IF EXISTS `c`.`s`.`mv`" in sql.statements
    assert "mv" not in sql.exists
    assert governance_failures(results)          # the job still fails
    assert _create_row(results, "c.s.t").action == "CREATE_OR_SKIP"  # table untouched


def test_pre_existing_mv_is_marked_not_dropped(tmp_path: Path):
    sql = _StatefulTarget(existing={"mv"}, fail_mv_tag=True)
    results = _run_migrating_mvs(tmp_path, sql)
    mv = _create_row(results, "c.s.mv")
    assert mv.action == "GOVERNANCE_FAILED"
    assert mv.error_code == "PROTECTION_FAILED"
    assert not [s for s in sql.statements if s.startswith("DROP")]
    assert "mv" in sql.exists


def test_fresh_mv_dropped_when_catalog_abac_policy_fails(tmp_path: Path):
    sql = _StatefulTarget(fail_catalog_policy=True)
    results = _run_migrating_mvs(tmp_path, sql)
    assert _create_row(results, "c.s.t").action == "DROP_PROTECTION_FAILED"
    assert _create_row(results, "c.s.mv").action == "DROP_PROTECTION_FAILED"
    assert "DROP MATERIALIZED VIEW IF EXISTS `c`.`s`.`mv`" in sql.statements
    assert "DROP TABLE IF EXISTS `c`.`s`.`t`" in sql.statements


def test_fresh_mv_kept_when_governance_succeeds(tmp_path: Path):
    sql = _StatefulTarget()
    results = _run_migrating_mvs(tmp_path, sql)
    mv = _create_row(results, "c.s.mv")
    assert mv.action == "CREATE_OR_SKIP" and mv.status == "SUCCESS"
    assert not [s for s in sql.statements if s.startswith("DROP")]
    assert governance_failures(results) == []


def test_policy_on_the_mv_itself_fails_and_drops_the_fresh_mv(tmp_path: Path):
    """Known ordering: a policy attached ON an MV is created before the MV exists, so
    it fails; being unprotected, the fresh MV is then dropped (the job fails)."""
    sql = _StatefulTarget()
    results = _run_migrating_mvs(tmp_path, sql, mv_level_policy=True)
    assert _create_row(results, "c.s.mv").action == "DROP_PROTECTION_FAILED"
    assert governance_failures(results)


# ---------------------------------------------------------------------------
# Report: governance rows of a report-only object read "Skipped (object not migrated)"
# ---------------------------------------------------------------------------


def test_report_marks_report_only_governance_rows_not_migrated(tmp_path: Path):
    from openpyxl import load_workbook

    from uc_sync.report import build_report

    tags = {"ssn": {"pii": "ssn"}}
    objects = [
        {"object_type": "TABLE", "full_name": "c.s.t",
         "definition": {"column_tags": tags}},
        {"object_type": "MATERIALIZED_VIEW", "full_name": "c.s.mv",
         "definition": {"column_tags": tags}},
        {"object_type": "ABAC_POLICY", "full_name": "c#policy:pii_mask",
         "definition": {"policy_name": "pii_mask", "on_securable_type": "CATALOG",
                        "on_securable": "c", "match_columns": ["has_tag('pii')"]}},
        {"object_type": "ABAC_POLICY", "full_name": "c.s.mv#policy:mv_mask",
         "definition": {"policy_name": "mv_mask", "on_securable_type": "TABLE",
                        "on_securable": "c.s.mv", "match_columns": ["has_tag('pii')"]}},
    ]
    import_results = [
        {"object_type": "TABLE", "target_full_name": "c.s.t", "full_name": "c.s.t",
         "status": "SUCCESS", "action": "CREATE_OR_SKIP"},
        {"object_type": "MATERIALIZED_VIEW", "target_full_name": "c.s.mv",
         "full_name": "c.s.mv", "status": "SUCCESS", "action": "REPORT_ONLY"},
        {"object_type": "TABLE", "target_full_name": "c.s.t", "full_name": "c.s.t",
         "status": "SUCCESS", "action": "APPLY_TAGS", "policies_path": "tags/x.sql"},
        {"object_type": "ABAC_POLICY", "target_full_name": "c#policy:pii_mask",
         "full_name": "c#policy:pii_mask", "status": "SUCCESS",
         "action": "CREATE_POLICY", "policies_path": "abac/x.sql"},
    ]
    out = tmp_path / "import.xlsx"
    build_report(objects, str(out), stage="IMPORT", import_results=import_results,
                 run_id="r1")
    wb = load_workbook(out)

    def rows(sheet):
        return [r for r in wb[sheet].iter_rows(min_row=2, values_only=True)]

    skipped = "Skipped (object not migrated)"
    tag_status = {r[0]: r[-1] for r in rows("Tags Applied")}
    assert tag_status == {"c.s.t": "APPLIED", "c.s.mv": skipped}

    abac_status = {r[0]: r[-1] for r in rows("ABAC Policies")}
    assert abac_status["mv_mask"] == skipped
    assert abac_status["pii_mask"] != skipped          # catalog policy: its own status

    matched = {(r[0], r[5]): r[-1] for r in rows("Policy Matched Columns")}
    assert matched[("pii_mask", "c.s.mv")] == skipped
    assert matched[("pii_mask", "c.s.t")] == abac_status["pii_mask"]
    assert matched[("mv_mask", "c.s.mv")] == skipped
