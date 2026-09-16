"""Bug A1/A2 + A3/A4: classic mask / row filter on a PRE-EXISTING column, and the
honest labels for a column drop / type change on a CHANGED table.

Root cause (fixed here): masks/row filters ride inline in SHOW CREATE, so adding one to
an already-migrated column lands only in ddl_hash → the table is CHANGED but
SKIP_EXISTING, and the active engine never re-issued the standalone SET MASK — it
falsely read "Updated". Phase 1c (``_apply_policies``) now re-applies the exported
``policies/`` dir to such tables, fail-closed.
"""

from __future__ import annotations

import json
from pathlib import Path

from uc_sync.fingerprints import (
    ddl_fingerprint,
    governance_fingerprint,
    grant_fingerprint_set,
)
from uc_sync.package_import import PackageImportEngine


class MaskSql:
    """Fake executor: CREATE TABLE reports already-exists (pre-existing target), DESCRIBE
    returns the given columns, SET MASK / SET ROW FILTER optionally fails."""

    def __init__(self, cols, *, fail_policy=False, describe_types=None):
        self.statements: list[str] = []
        self.cols = list(cols)
        self.fail_policy = fail_policy
        self.describe_types = describe_types or {}

    def execute(self, sql):
        self.statements.append(sql)
        u = sql.strip().upper()
        if u.startswith("CREATE TABLE") or u.startswith("CREATE EXTERNAL TABLE"):
            raise RuntimeError("[TABLE_OR_VIEW_ALREADY_EXISTS] exists")
        if u.startswith("DESCRIBE"):
            return [[c, self.describe_types.get(c, "string"), ""] for c in self.cols]
        if ("SET MASK" in u or "SET ROW FILTER" in u) and self.fail_policy:
            raise RuntimeError("[NO_SUCH_FUNCTION] mask/filter function missing")
        return None


def _write(root: Path, rel: str, body: str) -> None:
    p = root / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(body, encoding="utf-8")


def _bundle(tmp_path: Path, table_def: dict, *, policies: str = "") -> Path:
    root = tmp_path / "migrated"
    _write(root, "ddl/CATALOG_c.sql", "CREATE CATALOG `c`;\n")
    _write(root, "ddl/SCHEMA_c__s.sql", "CREATE SCHEMA `c`.`s`;\n")
    _write(root, "ddl/TABLE_c__s__t.sql",
           "CREATE TABLE `c`.`s`.`t` (id INT, note STRING);\n")
    if policies:
        _write(root, "policies/TABLE_c__s__t.sql", policies)
    inventory = [
        {"object_type": "CATALOG", "full_name": "c", "target_full_name": "c",
         "tags": {}, "grants": [], "definition": {}},
        {"object_type": "SCHEMA", "full_name": "c.s", "target_full_name": "c.s",
         "tags": {}, "grants": [], "definition": {}},
        {"object_type": "TABLE", "full_name": "c.s.t", "target_full_name": "c.s.t",
         "tags": {}, "grants": [], "definition": table_def},
    ]
    _write(root, "inventory/objects.json", json.dumps(inventory))
    return root


def _baseline(inventory_rows: list[dict], *, changed_table_hash: str = "OLD") -> dict:
    """Real fingerprints for catalog/schema (→ UNCHANGED, skipped), a stale ddl_hash for
    the table (→ CHANGED). last_action clean so nothing is force-re-attempted."""
    base = {}
    for row in inventory_rows:
        base[row["full_name"]] = {
            "object_type": row["object_type"],
            "ddl_hash": (
                changed_table_hash if row["object_type"] == "TABLE"
                else ddl_fingerprint(row)
            ),
            "governance_hash": governance_fingerprint(row),
            "grants": grant_fingerprint_set(row),
            "last_action": "created",
        }
    return base


def _run(root: Path, sql) -> list:
    inv = json.loads((root / "inventory" / "objects.json").read_text())
    return PackageImportEngine(
        str(root), sql, dry_run=False, prior_state=_baseline(inv)
    ).run()


# --- A1: classic mask on a pre-existing column ------------------------------


def test_mask_on_existing_column_is_applied_on_incremental(tmp_path):
    table_def = {
        "columns": [{"name": "id", "type_text": "int"},
                    {"name": "note", "type_text": "string"}],
        "column_masks": [{"column_name": "note", "function_name": "c.s.m",
                          "using_column_names": []}],
    }
    policies = "ALTER TABLE `c`.`s`.`t` ALTER COLUMN `note` SET MASK `c`.`s`.`m`;\n"
    root = _bundle(tmp_path, table_def, policies=policies)
    sql = MaskSql(cols=["id", "note"])
    results = _run(root, sql)

    # The mask on the pre-existing `note` column was actually re-applied.
    assert any("SET MASK" in s and "`note`" in s for s in sql.statements), sql.statements
    # A dedicated APPLY_POLICY result records success; the table reads Updated (changed
    # AND the mask landed), never a false success with no mask.
    pol = next(r for r in results if r.action == "APPLY_POLICY")
    assert pol.status == "SUCCESS"
    tbl = next(r for r in results if r.object_type == "TABLE" and r.action != "APPLY_POLICY")
    assert tbl.status in ("SUCCESS",) and tbl.delta_action == "CHANGED"


def test_mask_apply_failure_is_fail_closed_not_false_updated(tmp_path):
    """The core fix: a mask that cannot be applied must NEVER read Updated/success —
    the table is marked FAILURE (fail-closed), not silently left unprotected."""
    table_def = {
        "columns": [{"name": "id", "type_text": "int"},
                    {"name": "note", "type_text": "string"}],
        "column_masks": [{"column_name": "note", "function_name": "c.s.m",
                          "using_column_names": []}],
    }
    policies = "ALTER TABLE `c`.`s`.`t` ALTER COLUMN `note` SET MASK `c`.`s`.`m`;\n"
    root = _bundle(tmp_path, table_def, policies=policies)
    sql = MaskSql(cols=["id", "note"], fail_policy=True)
    results = _run(root, sql)

    pol = next(r for r in results if r.action == "APPLY_POLICY")
    assert pol.status == "FAILURE" and pol.error_code == "PROTECTION_FAILED"
    # The pre-existing table's own row is flipped to FAILURE in place — never Updated.
    tbl = next(
        r for r in results
        if r.object_type == "TABLE" and r.action != "APPLY_POLICY"
    )
    assert tbl.status == "FAILURE"
    assert tbl.error_code == "PROTECTION_FAILED"


def test_row_filter_on_existing_table_is_applied(tmp_path):
    table_def = {
        "columns": [{"name": "id", "type_text": "int"},
                    {"name": "dept", "type_text": "string"}],
        "row_filter": {"function_name": "c.s.rf", "input_column_names": ["dept"]},
    }
    policies = "ALTER TABLE `c`.`s`.`t` SET ROW FILTER `c`.`s`.`rf` ON (`dept`);\n"
    root = _bundle(tmp_path, table_def, policies=policies)
    sql = MaskSql(cols=["id", "dept"])
    results = _run(root, sql)
    assert any("SET ROW FILTER" in s for s in sql.statements), sql.statements
    assert next(r for r in results if r.action == "APPLY_POLICY").status == "SUCCESS"


def test_full_run_does_not_run_policy_phase(tmp_path):
    """On a full run masks ride inline in CREATE TABLE — no separate policy pass."""
    table_def = {
        "columns": [{"name": "id", "type_text": "int"},
                    {"name": "note", "type_text": "string"}],
        "column_masks": [{"column_name": "note", "function_name": "c.s.m",
                          "using_column_names": []}],
    }
    policies = "ALTER TABLE `c`.`s`.`t` ALTER COLUMN `note` SET MASK `c`.`s`.`m`;\n"
    root = _bundle(tmp_path, table_def, policies=policies)
    sql = MaskSql(cols=["id", "note"])
    results = PackageImportEngine(str(root), sql, dry_run=False).run()  # no baseline
    assert not any(r.action == "APPLY_POLICY" for r in results)


# --- A3/A4: column drop / type change read Skipped-with-comment -------------


def test_column_dropped_on_source_reads_skipped_with_comment(tmp_path):
    # Source has only `id`; target still has `id`, `note` → `note` dropped on source.
    table_def = {"columns": [{"name": "id", "type_text": "int"}]}
    root = _bundle(tmp_path, table_def)
    sql = MaskSql(cols=["id", "note"])  # target still has the dropped column
    results = _run(root, sql)
    tbl = next(r for r in results if r.object_type == "TABLE" and r.action != "APPLY_POLICY")
    assert tbl.delta_action == "CHANGED_SKIPPED"
    assert "deleted on source" in tbl.message
    # No destructive ALTER / DROP ran.
    assert not any("DROP" in s.upper() and "COLUMN" in s.upper() for s in sql.statements)


def test_column_type_change_reads_skipped_with_comment(tmp_path):
    # Same column set, but `id` is bigint on source vs int on target → type change.
    table_def = {"columns": [{"name": "id", "type_text": "bigint"},
                             {"name": "note", "type_text": "string"}]}
    root = _bundle(tmp_path, table_def)
    sql = MaskSql(cols=["id", "note"], describe_types={"id": "int", "note": "string"})
    results = _run(root, sql)
    tbl = next(r for r in results if r.object_type == "TABLE" and r.action != "APPLY_POLICY")
    assert tbl.delta_action == "CHANGED_SKIPPED"
    assert "type changed on source" in tbl.message
    assert not any("ALTER COLUMN" in s.upper() and "TYPE" in s.upper()
                   for s in sql.statements)


def test_changed_skipped_reads_as_skipped_label_and_reason_in_report(tmp_path):
    """End-to-end: the CHANGED_SKIPPED delta action renders a Skipped variant WITH the
    reason on the per-type sheet, and the shared vocab maps it to 'skipped'."""
    from openpyxl import load_workbook
    from uc_sync import vocab
    from uc_sync.report import build_report

    assert vocab.status_key({"status": "SUCCESS", "action": "SKIP_EXISTING",
                             "delta_action": "CHANGED_SKIPPED"}) == "skipped"

    objects = [{"object_type": "TABLE", "full_name": "c.s.t", "owner": "me",
                "tags": {}, "grants": []}]
    import_results = [{
        "object_type": "TABLE", "target_full_name": "c.s.t", "full_name": "c.s.t",
        "status": "SUCCESS", "action": "SKIP_EXISTING", "delta_action": "CHANGED_SKIPPED",
        "message": "column deleted on source (note); not dropped on target (non-destructive)",
    }]
    out = tmp_path / "r.xlsx"
    build_report(objects, str(out), stage="IMPORT", import_results=import_results,
                 run_id="r1")
    row = {r[0]: r[-1] for r in load_workbook(out)["Tables"].iter_rows(values_only=True)}
    assert row["c.s.t"].startswith("Skipped —")
    assert "deleted on source" in row["c.s.t"]
