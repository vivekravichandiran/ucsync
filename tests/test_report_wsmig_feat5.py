"""FEAT-5: the report adopts the workspace-migration (wsmig) vocabulary + layout, and
supersedes #12 — a report-only object reads as a Skipped variant, counted OUTSIDE
success, never "SUCCESS (REPORT_ONLY)"."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from uc_sync.report import build_report


def _objects():
    return [
        {"object_type": "CATALOG", "full_name": "c", "target_full_name": "c",
         "tags": {}, "grants": [], "definition": {}},
        {"object_type": "TABLE", "full_name": "c.s.t", "target_full_name": "c.s.t",
         "tags": {}, "grants": [], "definition": {"columns": [{"name": "id"}]}},
        # Two report-only objects (never migrated).
        {"object_type": "STREAMING_TABLE", "full_name": "c.s.st",
         "target_full_name": "c.s.st", "tags": {}, "grants": [],
         "definition": {"in_scope_for_migration": False}},
        {"object_type": "MODEL", "full_name": "c.s.m", "target_full_name": "c.s.m",
         "tags": {}, "grants": [], "definition": {"in_scope_for_migration": False}},
    ]


def _import_results():
    return [
        {"object_type": "CATALOG", "full_name": "c", "target_full_name": "c",
         "source_full_name": "c", "status": "SUCCESS", "action": "CREATE_OR_SKIP",
         "message": ""},
        {"object_type": "TABLE", "full_name": "c.s.t", "target_full_name": "c.s.t",
         "source_full_name": "c.s.t", "status": "SUCCESS", "action": "CREATE_OR_SKIP",
         "message": ""},
    ]


def test_report_only_reads_skipped_and_counted_outside_success(tmp_path: Path):
    out = tmp_path / "r.xlsx"
    build_report(_objects(), str(out), stage="IMPORT",
                 import_results=_import_results(), workspace_url="https://ws.example")
    wb = load_workbook(out)

    # No cell anywhere reads the misleading "SUCCESS (REPORT_ONLY)".
    for sheet in wb.sheetnames:
        for row in wb[sheet].iter_rows(values_only=True):
            for c in row:
                assert "SUCCESS (REPORT_ONLY)" not in str(c or "")

    summary = [tuple(r) for r in wb["Summary"].iter_rows(values_only=True)]
    flat = {r[0]: r[1] for r in summary if r and len(r) >= 2}
    # 2 created (catalog + table), 2 report-only skipped (streaming table + model).
    assert flat.get("Created") == 2
    assert flat.get("Skipped (no target object)") == 2
    assert flat.get("applied (created/updated/adopted)") == 2
    assert flat.get("skipped (incl. report-only)") == 2
    # A TOTAL row is present and covers every counted object.
    assert flat.get("TOTAL") == 4


def test_summary_title_bar_has_workspace_and_timestamp(tmp_path: Path):
    out = tmp_path / "r.xlsx"
    build_report(_objects(), str(out), stage="IMPORT",
                 import_results=_import_results(), workspace_url="https://ws.example")
    summary = [tuple(r) for r in load_workbook(out)["Summary"].iter_rows(values_only=True)]
    flat = {r[0]: r[1] for r in summary if r and len(r) >= 2}
    assert flat.get("workspace") == "https://ws.example"
    assert "generated" in {r[0] for r in summary if r}


def test_failures_listed_first_in_rollup(tmp_path: Path):
    objs = _objects() + [
        {"object_type": "TABLE", "full_name": "c.s.bad", "target_full_name": "c.s.bad",
         "tags": {}, "grants": [], "definition": {}},
    ]
    results = _import_results() + [
        {"object_type": "TABLE", "full_name": "c.s.bad", "target_full_name": "c.s.bad",
         "source_full_name": "c.s.bad", "status": "FAILURE", "action": "CREATE_OR_SKIP",
         "error_code": "PROTECTION_FAILED", "message": "dropped"},
    ]
    out = tmp_path / "r.xlsx"
    build_report(objs, str(out), stage="IMPORT", import_results=results)
    summary = [r[0] for r in load_workbook(out)["Summary"].iter_rows(values_only=True) if r]
    # FAILED appears before Created in the roll-up (failures first).
    roll = summary[summary.index("Outcome roll-up"):]
    assert roll.index("FAILED") < roll.index("Created")


def test_status_style_covers_all_summary_order():
    from uc_sync.report import _STATUS_STYLE, _SUMMARY_ORDER
    for key in _SUMMARY_ORDER:
        assert key in _STATUS_STYLE
